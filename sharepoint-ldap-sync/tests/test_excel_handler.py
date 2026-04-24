"""
Unit tests for the ExcelHandler class.

These tests run entirely in-memory (no SharePoint, no LDAP) using openpyxl
to construct workbooks and verify round-trip behaviour.
"""

from __future__ import annotations

import io
import sys
import os

# Ensure the package root is on the path so imports resolve without installing.
sys.path.insert(0, os.path.join(os.path.dirname(__file__), ".."))

import openpyxl
import pytest

from config import ExcelConfig
from excel_handler import ExcelHandler


# ---------------------------------------------------------------------------
# Fixtures
# ---------------------------------------------------------------------------

def _make_workbook(rows: list[list]) -> bytes:
    """Create a minimal workbook and return it as bytes."""
    wb = openpyxl.Workbook()
    ws = wb.active
    for r_idx, row in enumerate(rows, start=1):
        for c_idx, val in enumerate(row, start=1):
            ws.cell(row=r_idx, column=c_idx, value=val)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


DEFAULT_CFG = ExcelConfig(
    username_column="A",
    header_row=1,
    data_start_row=2,
)


# ---------------------------------------------------------------------------
# Tests
# ---------------------------------------------------------------------------

class TestLoadAndIterUsernames:
    def test_basic_usernames(self):
        data = _make_workbook([
            ["Username"],
            ["jsmith"],
            ["mjones"],
            ["bwilliams"],
        ])
        handler = ExcelHandler(DEFAULT_CFG)
        handler.load_from_bytes(data)
        result = handler.get_usernames()
        assert result == [(2, "jsmith"), (3, "mjones"), (4, "bwilliams")]

    def test_skips_empty_rows(self):
        data = _make_workbook([
            ["Username"],
            ["jsmith"],
            [None],
            ["mjones"],
        ])
        handler = ExcelHandler(DEFAULT_CFG)
        handler.load_from_bytes(data)
        result = handler.get_usernames()
        assert result == [(2, "jsmith"), (4, "mjones")]

    def test_strips_whitespace(self):
        data = _make_workbook([
            ["Username"],
            ["  jsmith  "],
        ])
        handler = ExcelHandler(DEFAULT_CFG)
        handler.load_from_bytes(data)
        assert handler.get_usernames() == [(2, "jsmith")]

    def test_empty_sheet(self):
        data = _make_workbook([["Username"]])
        handler = ExcelHandler(DEFAULT_CFG)
        handler.load_from_bytes(data)
        assert handler.get_usernames() == []

    def test_no_workbook_raises(self):
        handler = ExcelHandler(DEFAULT_CFG)
        with pytest.raises(RuntimeError):
            list(handler.iter_usernames())


class TestWriteUserData:
    def test_new_columns_are_created(self):
        data = _make_workbook([["Username"], ["jsmith"]])
        handler = ExcelHandler(DEFAULT_CFG)
        handler.load_from_bytes(data)
        handler.write_user_data(2, {"mail": "jsmith@example.com", "displayName": "John Smith"})

        # Reload to verify persistence
        out = handler.to_bytes()
        wb2 = openpyxl.load_workbook(io.BytesIO(out))
        ws2 = wb2.active

        # Header row must contain the attribute names
        headers = {ws2.cell(1, c).value for c in range(1, ws2.max_column + 1)}
        assert "mail" in headers
        assert "displayName" in headers

    def test_values_written_to_correct_row(self):
        data = _make_workbook([["Username"], ["jsmith"], ["mjones"]])
        handler = ExcelHandler(DEFAULT_CFG)
        handler.load_from_bytes(data)
        handler.write_user_data(2, {"mail": "jsmith@example.com"})
        handler.write_user_data(3, {"mail": "mjones@example.com"})

        out = handler.to_bytes()
        wb2 = openpyxl.load_workbook(io.BytesIO(out))
        ws2 = wb2.active

        # Find the mail column
        mail_col = None
        for cell in ws2[1]:
            if cell.value == "mail":
                mail_col = cell.column
        assert mail_col is not None

        assert ws2.cell(2, mail_col).value == "jsmith@example.com"
        assert ws2.cell(3, mail_col).value == "mjones@example.com"

    def test_existing_column_reused(self):
        data = _make_workbook([
            ["Username", "mail"],
            ["jsmith", "old@example.com"],
        ])
        handler = ExcelHandler(DEFAULT_CFG)
        handler.load_from_bytes(data)
        handler.write_user_data(2, {"mail": "new@example.com"})

        out = handler.to_bytes()
        wb2 = openpyxl.load_workbook(io.BytesIO(out))
        ws2 = wb2.active

        # The mail column should still be column B (index 2)
        assert ws2.cell(2, 2).value == "new@example.com"
        # And there should be no duplicate header
        headers = [ws2.cell(1, c).value for c in range(1, ws2.max_column + 1)]
        assert headers.count("mail") == 1


class TestToBytes:
    def test_round_trip(self):
        data = _make_workbook([["Username"], ["jsmith"]])
        handler = ExcelHandler(DEFAULT_CFG)
        handler.load_from_bytes(data)
        out = handler.to_bytes()
        assert isinstance(out, bytes)
        assert len(out) > 0

    def test_raises_without_workbook(self):
        handler = ExcelHandler(DEFAULT_CFG)
        with pytest.raises(RuntimeError):
            handler.to_bytes()


class TestCustomConfig:
    def test_different_username_column(self):
        cfg = ExcelConfig(username_column="B", header_row=1, data_start_row=2)
        data = _make_workbook([
            ["ID", "Username"],
            [1, "jsmith"],
            [2, "mjones"],
        ])
        handler = ExcelHandler(cfg)
        handler.load_from_bytes(data)
        result = handler.get_usernames()
        assert result == [(2, "jsmith"), (3, "mjones")]
